from pathlib import Path
from typing import Optional

SUPPORTED_EXTENSIONS = {".md", ".txt", ".pdf", ".docx", ".xlsx", ".csv"}


def parse_file(file_path: str, original_filename: Optional[str] = None) -> dict:
    path = Path(file_path)
    if not path.exists():
        return {"ok": False, "error": f"文件不存在: {file_path}"}

    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        return {"ok": False, "error": f"不支持的文件类型: {ext}"}

    fname = original_filename or path.name

    try:
        if ext in (".md", ".txt"):
            return _parse_text(path, fname)
        elif ext == ".pdf":
            return _parse_pdf(path, fname)
        elif ext == ".docx":
            return _parse_docx(path, fname)
        elif ext in (".xlsx", ".csv"):
            return _parse_tabular(path, fname, ext)
        else:
            return {"ok": False, "error": f"未实现的解析器: {ext}"}
    except Exception as e:
        return {"ok": False, "error": f"解析失败: {str(e)}"}


def _parse_text(path: Path, filename: str) -> dict:
    content = path.read_text(encoding="utf-8", errors="replace")
    return {
        "ok": True,
        "content": content,
        "source_type": "markdown" if path.suffix == ".md" else "txt",
        "title": path.stem,
        "original_filename": filename,
        "char_count": len(content),
    }


def _parse_pdf(path: Path, filename: str) -> dict:
    try:
        import fitz
    except ImportError:
        return {"ok": False, "error": "PDF解析需要安装PyMuPDF: pip install PyMuPDF"}

    doc = fitz.open(str(path))
    pages = []
    for page in doc:
        pages.append(page.get_text())
    doc.close()

    content = "\n\n".join(pages)
    return {
        "ok": True,
        "content": content,
        "source_type": "pdf",
        "title": path.stem,
        "original_filename": filename,
        "char_count": len(content),
    }


def _parse_docx(path: Path, filename: str) -> dict:
    try:
        from docx import Document
    except ImportError:
        return {
            "ok": False,
            "error": "DOCX解析需要安装python-docx: pip install python-docx",
        }

    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    content = "\n\n".join(paragraphs)

    return {
        "ok": True,
        "content": content,
        "source_type": "docx",
        "title": path.stem,
        "original_filename": filename,
        "char_count": len(content),
    }


def _parse_tabular(path: Path, filename: str, ext: str) -> dict:
    try:
        import openpyxl
    except ImportError:
        return {"ok": False, "error": "表格解析需要安装openpyxl: pip install openpyxl"}

    rows = []
    if ext == ".xlsx":
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                rows.append(" | ".join(cells))
        wb.close()
    elif ext == ".csv":
        import csv

        with open(str(path), "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            for row in reader:
                rows.append(" | ".join(row))

    content = "\n".join(rows)
    return {
        "ok": True,
        "content": content,
        "source_type": ext.lstrip("."),
        "title": path.stem,
        "original_filename": filename,
        "char_count": len(content),
    }
